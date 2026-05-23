from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "outputs"
OUT_DIR.mkdir(exist_ok=True)

PNG_PATH = OUT_DIR / "hand_bridge_poster_report.png"
DOCX_PATH = OUT_DIR / "hand_bridge_poster_report.docx"

W, H = 1080, 1600
BG = (246, 248, 251)
BLACK = (18, 18, 18)
LABEL = (184, 184, 184)
ORANGE = (255, 169, 54)
ORANGE_DARK = (230, 122, 0)
WHITE = (255, 255, 255)
BLUE = (224, 239, 255)
GREEN = (227, 245, 233)
YELLOW = (255, 246, 216)

FONT_REG = "/System/Library/Fonts/AppleSDGothicNeo.ttc"
FONT_BOLD = "/System/Library/Fonts/AppleSDGothicNeo.ttc"


def font(size, bold=False):
    # AppleSDGothicNeo.ttc exposes bold faces through font index on macOS.
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size, index=1 if bold else 0)


F_TITLE = font(43)
F_SECTION = font(24, True)
F_BODY = font(17)
F_BODY_B = font(17, True)
F_SMALL = font(15)
F_SMALL_B = font(15, True)
F_TINY = font(13)


def rounded(draw, xy, radius=24, outline=BLACK, width=2, fill=None):
    draw.rounded_rectangle(xy, radius=radius, outline=outline, width=width, fill=fill)


def label(draw, x, y, text, width):
    draw.rounded_rectangle((x, y, x + width, y + 32), radius=7, fill=LABEL)
    draw.ellipse((x - 16, y - 8, x + 22, y + 30), fill=ORANGE, outline=ORANGE_DARK, width=3)
    draw.text((x + 34, y + 3), text, fill=BLACK, font=F_SECTION)


def draw_wrapped(draw, text, x, y, max_width, line_h=29, fill=BLACK, fnt=F_BODY, bullet_indent=0):
    cur = y
    for raw in text.split("\n"):
        if raw == "":
            cur += line_h // 2
            continue
        words = raw.split(" ")
        line = ""
        for word in words:
            test = word if not line else line + " " + word
            if draw.textlength(test, font=fnt) <= max_width:
                line = test
            else:
                draw.text((x + bullet_indent, cur), line, fill=fill, font=fnt)
                cur += line_h
                line = word
        if line:
            draw.text((x + bullet_indent, cur), line, fill=fill, font=fnt)
            cur += line_h
    return cur


def bullet(draw, text, x, y, max_width, fnt=F_BODY, line_h=27):
    draw.text((x, y), "·", fill=BLACK, font=fnt)
    return draw_wrapped(draw, text, x + 18, y, max_width - 18, line_h=line_h, fnt=fnt)


def mini_badge(draw, x, y, text, fill):
    draw.rounded_rectangle((x, y, x + 116, y + 28), radius=8, fill=fill, outline=(190, 190, 190), width=1)
    tw = draw.textlength(text, font=F_TINY)
    draw.text((x + (116 - tw) / 2, y + 6), text, fill=BLACK, font=F_TINY)


def draw_flow(draw, x, y):
    items = [
        ("Webcam", BLUE),
        ("MediaPipe\n관절 추출", GREEN),
        ("LSTM\n동작 분류", YELLOW),
        ("Korean\nText", BLUE),
    ]
    box_w, box_h, gap = 90, 55, 16
    for i, (txt, fill) in enumerate(items):
        bx = x + i * (box_w + gap)
        draw.rounded_rectangle((bx, y, bx + box_w, y + box_h), radius=9, fill=fill, outline=BLACK, width=1)
        lines = txt.split("\n")
        for j, line in enumerate(lines):
            tw = draw.textlength(line, font=F_TINY)
            draw.text((bx + (box_w - tw) / 2, y + 12 + j * 17), line, fill=BLACK, font=F_TINY)
        if i < len(items) - 1:
            ax = bx + box_w + 2
            ay = y + box_h / 2
            draw.line((ax, ay, ax + gap - 5, ay), fill=BLACK, width=2)
            draw.polygon([(ax + gap - 5, ay), (ax + gap - 12, ay - 5), (ax + gap - 12, ay + 5)], fill=BLACK)


def draw_ui_mock(draw, x, y, w, h):
    draw.rounded_rectangle((x, y, x + w, y + h), radius=12, fill=(31, 33, 38), outline=BLACK, width=1)
    draw.rectangle((x + 12, y + 12, x + w - 12, y + 104), fill=(54, 58, 66))
    for i in range(9):
        px = x + 45 + i * 32
        py = y + 28 + (i % 3) * 12
        draw.ellipse((px, py, px + 6, py + 6), fill=ORANGE)
        if i > 0:
            draw.line((px - 26, py - 8, px, py), fill=(255, 214, 130), width=1)
    draw.rounded_rectangle((x + 14, y + 118, x + w - 14, y + 151), radius=7, fill=(238, 238, 238))
    draw.text((x + 25, y + 125), "현재 수어: 감사합니다  |  신뢰도 87%", fill=BLACK, font=F_TINY)


img = Image.new("RGB", (W, H), BG)
draw = ImageDraw.Draw(img)

# Title block
rounded(draw, (22, 6, 1052, 188), radius=34, outline=(55, 55, 55), width=1, fill=WHITE)
title = "MediaPipe를 이용한 수어 번역기"
tw = draw.textlength(title, font=F_TITLE)
draw.text(((W - tw) / 2, 53), title, fill=BLACK, font=F_TITLE)

# Layout boxes
boxes = {
    "purpose": (22, 235, 533, 420),
    "theory": (22, 448, 533, 808),
    "process": (22, 852, 533, 1558),
    "result": (560, 235, 1052, 1077),
    "conclusion": (560, 1133, 1052, 1348),
    "refs": (560, 1400, 1052, 1560),
}
for key, xy in boxes.items():
    rounded(draw, xy, radius=26, outline=BLACK, width=2, fill=None)

label(draw, 54, 220, "연구목적과 필요성", 230)
label(draw, 54, 433, "이론적 고찰", 225)
label(draw, 54, 838, "연구 과정", 215)
label(draw, 594, 220, "연 구 결 과", 225)
label(draw, 594, 1118, "결론 및 기대효과", 220)
label(draw, 594, 1386, "참 고 문 헌", 215)

# Purpose
y = 280
draw.text((75, y), "본 연구는", fill=BLACK, font=F_BODY_B)
y += 36
purpose = (
    "청각장애인과 비수어 사용자 사이의 의사소통 장벽을 줄이기 위해 "
    "웹캠으로 손동작을 인식하고 한국어 문장으로 보여주는 실시간 수어 번역기를 제작하는 것을 목표로 한다."
)
draw_wrapped(draw, purpose, 75, y, 405, line_h=29)

# Theory
y = 532
draw.text((52, y), "1.", fill=BLACK, font=F_BODY_B)
draw.text((87, y), "MediaPipe", fill=BLACK, font=F_BODY_B)
y = bullet(draw, "Google에서 개발한 프레임워크로, 별도 GPU 없이도 손의 21개 관절 좌표를 실시간으로 추적한다.", 86, y + 36, 395)
y += 18
draw.text((52, y), "2.", fill=BLACK, font=F_BODY_B)
draw.text((87, y), "TensorFlow / Keras", fill=BLACK, font=F_BODY_B)
y = bullet(draw, "추출한 좌표 데이터를 학습시켜 수어 동작을 분류하는 딥러닝 모델을 만들기 위해 사용하였다.", 86, y + 36, 395)
y += 18
draw.text((52, y), "3.", fill=BLACK, font=F_BODY_B)
draw.text((87, y), "OpenCV · NumPy · LSTM", fill=BLACK, font=F_BODY_B)
y = bullet(draw, "영상 처리·좌표 계산·동작 흐름 분석을 담당한다.", 86, y + 36, 395)

# Process
y = 895
draw.text((47, y), "가. Python", fill=BLACK, font=F_BODY_B)
y += 36
draw.text((52, y), "1) Python을 사용한 이유", fill=BLACK, font=F_SMALL_B)
y = bullet(draw, "AI, 영상 처리, GUI 라이브러리가 풍부해 짧은 시간 안에 실시간 프로그램을 구현하기 좋다.", 55, y + 30, 415, fnt=F_SMALL, line_h=25)
y += 18
draw.text((47, y), "나. 기능", fill=BLACK, font=F_BODY_B)
y += 34
steps = [
    ("1) 수어 번역기의 목표", "수어 동작을 카메라로 입력받아 한국어 텍스트로 변환한다."),
    ("2) 관절 포인트 추출", "MediaPipe로 양손의 x, y, z 좌표를 추출하고 30프레임 단위로 묶는다."),
    ("3) 데이터 저장", "각 단어별 좌표 데이터를 .npy 파일로 저장해 학습 자료로 사용한다."),
    ("4) 모델 학습", "TensorFlow 기반 LSTM 모델이 시간 순서가 있는 손동작 패턴을 학습한다."),
    ("5) 실시간 번역", "예측값이 일정 신뢰도 이상이면 현재 수어와 번역 문장을 화면에 출력한다."),
]
for head, body in steps:
    draw.text((52, y), head, fill=BLACK, font=F_SMALL_B)
    y = bullet(draw, body, 68, y + 25, 390, fnt=F_SMALL, line_h=24)
    y += 12

mini_badge(draw, 68, 1458, "로컬 실행", GREEN)
mini_badge(draw, 198, 1458, "GPU 불필요", YELLOW)
mini_badge(draw, 328, 1458, "확장 가능", BLUE)

# Result
y = 272
draw.text((588, y), "가. 제작", fill=BLACK, font=F_BODY_B)
y += 35
draw_flow(draw, 606, y)
y += 82
draw_wrapped(draw, "웹캠 입력 → 손 관절 좌표 추출 → 30프레임 분석 → 수어 단어 예측 → 한글 출력 구조로 제작하였다.", 590, y, 405, line_h=27, fnt=F_SMALL)

y += 92
draw.text((588, y), "나. 실행화면", fill=BLACK, font=F_BODY_B)
draw_ui_mock(draw, 605, y + 38, 360, 170)

y += 258
draw.text((588, y), "다. 결과", fill=BLACK, font=F_BODY_B)
y += 38
result = (
    "카메라 화면에서 손의 움직임을 추적하고, 학습된 단어가 입력되면 번역 결과와 신뢰도를 표시하였다. "
    "현재 프로젝트는 '감사합니다', '만나서' 등 수어 데이터를 기반으로 동작 인식 흐름을 확인할 수 있다."
)
y = draw_wrapped(draw, result, 590, y, 405, line_h=28)
y += 22
for label_text, value, fill in [
    ("입력", "웹캠 영상", BLUE),
    ("처리", "관절 좌표 + LSTM", GREEN),
    ("출력", "한국어 텍스트", YELLOW),
]:
    draw.rounded_rectangle((590, y, 1000, y + 43), radius=9, fill=fill, outline=(155, 155, 155), width=1)
    draw.text((610, y + 10), label_text, fill=BLACK, font=F_SMALL_B)
    draw.text((695, y + 10), value, fill=BLACK, font=F_SMALL)
    y += 54

# Conclusion
y = 1177
conclusion = (
    "본 프로그램은 수어 동작을 좌표 데이터로 변환하고 AI 모델로 분류해 실시간 번역 가능성을 보여주었다. "
    "향후 더 많은 단어와 다양한 사용자 영상을 추가하면 인식률을 높이고, 교육·안내·접근성 보조 도구로 확장할 수 있다."
)
draw_wrapped(draw, conclusion, 590, y, 405, line_h=29)

# References
y = 1432
refs = [
    "MediaPipe 공식 문서",
    "TensorFlow / Keras 공식 문서",
    "OpenCV 공식 문서",
    "Hand-Bridge 프로젝트 소스 코드",
]
for r in refs:
    y = bullet(draw, r, 590, y, 395, fnt=F_SMALL, line_h=25)
    y += 4

img.save(PNG_PATH, quality=95)

doc = Document()
section = doc.sections[0]
section.page_width = Inches(6.75)
section.page_height = Inches(10)
section.top_margin = Inches(0)
section.bottom_margin = Inches(0)
section.left_margin = Inches(0)
section.right_margin = Inches(0)

p = doc.add_paragraph()
p.paragraph_format.space_after = 0
p.paragraph_format.space_before = 0
run = p.add_run()
run.add_picture(str(PNG_PATH), width=Inches(6.75), height=Inches(10))

doc.save(DOCX_PATH)
print(DOCX_PATH)
print(PNG_PATH)
